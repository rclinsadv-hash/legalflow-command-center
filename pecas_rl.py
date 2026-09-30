"""Gerador de peças no padrão do escritório (Raphael Correia Lins, OAB/PB 21.036).

Padrão extraído de peças reais do escritório: A4, margens 2,5/2,5/3,0/3,0 cm, Verdana 12,
justificado, recuo de 1,25 cm, entrelinha 1,5, títulos em azul-marinho e tabelas com
cabeçalho azul-marinho.

Uso mínimo:
    from pecas_rl import Peca
    p = Peca(modelo=r"G:\\...\\modelo_com_timbrado.docx")   # modelo opcional: herda cabeçalho/rodapé
    p.enderecamento("Excelentíssimo(a) Senhor(a) Juiz(a) de Direito ...")
    p.processo("0000000-00.2026.8.15.0561")
    p.titulo_peca("Recurso Inominado")
    p.secao("I. DA TEMPESTIVIDADE")
    p.paragrafo("Texto com **negrito** inline.")
    p.fecho("Coremas/PB", "30 de setembro de 2026")
    p.salvar(r"G:\\...\\CLIENTES NOVOS\\Fulano\\recurso.docx")

Limites conhecidos: não cria notas de rodapé nem insere imagens/figuras. Para isso, edite o
.docx gerado ou use a skill de Word. Cabeçalho e rodapé (papel timbrado) só existem se você
passar `modelo=` com um .docx do escritório.
"""
from __future__ import annotations

import copy
import os
import re
import shutil
import subprocess

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONTE = "Verdana"
AZUL = "1F3864"
VERMELHO = "CC0000"
DOURADO = "D4AF37"
CINZA_CLARO = "F0F4F8"

ADVOGADO = "Raphael Correia Lins"
OAB = "OAB/PB nº 21.036"

CORPO_PT = 12
TABELA_PT = 10
CITACAO_PT = 10
RECUO_PRIMEIRA_LINHA = Cm(1.25)
RECUO_CITACAO = Cm(4.0)


def _rgb(hexa: str) -> RGBColor:
    return RGBColor.from_string(hexa)


def _fonte(run, tamanho=CORPO_PT, negrito=None, italico=None, cor=None):
    run.font.name = FONTE
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for atributo in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(atributo), FONTE)
    run.font.size = Pt(tamanho)
    if negrito is not None:
        run.font.bold = negrito
    if italico is not None:
        run.font.italic = italico
    if cor:
        run.font.color.rgb = _rgb(cor)


def _sombrear(celula, hexa: str):
    tc_pr = celula._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexa)
    tc_pr.append(shd)


def _bordas(tabela, hexa="BFBFBF"):
    tbl_pr = tabela._element.tblPr
    bordas = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{lado}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), hexa)
        bordas.append(el)
    tbl_pr.append(bordas)


class Peca:
    """Monta uma peça .docx no padrão do escritório."""

    def __init__(self, modelo: str | None = None):
        if modelo:
            self.doc = Document(modelo)
            self._limpar_corpo()
            self.tem_timbrado = True
        else:
            self.doc = Document()
            self._configurar_pagina()
            self.tem_timbrado = False
        self._configurar_estilo_normal()

    # -- preparação -----------------------------------------------------
    def _limpar_corpo(self):
        """Remove o conteúdo do modelo, mantendo seção (margens, cabeçalho e rodapé)."""
        corpo = self.doc.element.body
        for filho in list(corpo):
            if filho.tag != qn("w:sectPr"):
                corpo.remove(filho)

    def _configurar_pagina(self):
        secao = self.doc.sections[0]
        secao.page_width, secao.page_height = Cm(21.0), Cm(29.7)
        secao.top_margin = secao.bottom_margin = Cm(2.5)
        secao.left_margin = secao.right_margin = Cm(3.0)

    def _configurar_estilo_normal(self):
        normal = self.doc.styles["Normal"]
        normal.font.name = FONTE
        normal.font.size = Pt(CORPO_PT)
        rpr = normal.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        for atributo in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rfonts.set(qn(atributo), FONTE)
        fmt = normal.paragraph_format
        fmt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        fmt.line_spacing = 1.5
        fmt.space_before = Pt(5)
        fmt.space_after = Pt(5)

    # -- utilitários ----------------------------------------------------
    def _escrever(self, paragrafo, texto, tamanho=CORPO_PT, negrito=None, italico=None, cor=None):
        """Escreve texto com **negrito** inline."""
        for i, trecho in enumerate(re.split(r"\*\*", texto)):
            if not trecho:
                continue
            run = paragrafo.add_run(trecho)
            _fonte(run, tamanho, True if i % 2 else negrito, italico, cor)
        return paragrafo

    # -- blocos da peça -------------------------------------------------
    def enderecamento(self, texto: str):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        self._escrever(p, texto.upper(), negrito=True)
        return p

    def processo(self, numero: str, extra: list[str] | None = None):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        self._escrever(p, f"Processo nº {numero}", negrito=True)
        for linha in extra or []:
            q = self.doc.add_paragraph()
            q.alignment = WD_ALIGN_PARAGRAPH.LEFT
            self._escrever(q, linha)
        return p

    def titulo_peca(self, texto: str):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(12)
        self._escrever(p, texto.upper(), tamanho=14, negrito=True, cor=AZUL)
        return p

    def secao(self, texto: str, nivel: int = 1):
        """Título de seção. Numere você mesmo: 'I. DA TEMPESTIVIDADE', 'IV.1. Do dano moral'."""
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(12 if nivel == 1 else 8)
        self._escrever(p, texto.upper() if nivel == 1 else texto, tamanho=13 if nivel == 1 else 12,
                       negrito=True, cor=AZUL)
        return p

    def paragrafo(self, texto: str, recuo: bool = True):
        p = self.doc.add_paragraph()
        if recuo:
            p.paragraph_format.first_line_indent = RECUO_PRIMEIRA_LINHA
        self._escrever(p, texto)
        return p

    def citacao(self, texto: str):
        """Citação/ementa: recuo de 4 cm, 10 pt, entrelinha simples."""
        p = self.doc.add_paragraph()
        p.paragraph_format.left_indent = RECUO_CITACAO
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        self._escrever(p, texto, tamanho=CITACAO_PT)
        return p

    def alerta(self, texto: str):
        """Destaque em vermelho (contradição, ponto crítico). Use com parcimônia."""
        p = self.doc.add_paragraph()
        self._escrever(p, texto, negrito=True, cor=VERMELHO)
        return p

    def pedidos(self, itens: list[str], estilo: str = "letras"):
        """Lista de pedidos: a), b), c) ou 1., 2., 3."""
        for i, item in enumerate(itens):
            marcador = f"{chr(ord('a') + i)})" if estilo == "letras" else f"{i + 1}."
            p = self.doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1.25)
            p.paragraph_format.first_line_indent = Cm(-0.75)
            self._escrever(p, f"{marcador} {item}")
        return self.doc

    def tabela(self, cabecalho: list[str], linhas: list[list[str]], larguras_cm: list[float] | None = None):
        """Tabela com cabeçalho azul-marinho e texto branco, corpo em 10 pt."""
        tabela = self.doc.add_table(rows=1, cols=len(cabecalho))
        tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
        _bordas(tabela)
        for i, titulo in enumerate(cabecalho):
            cel = tabela.rows[0].cells[i]
            _sombrear(cel, AZUL)
            par = cel.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.line_spacing = 1.0
            self._escrever(par, titulo, tamanho=TABELA_PT, negrito=True, cor="FFFFFF")
        for n, linha in enumerate(linhas):
            cels = tabela.add_row().cells
            for i, valor in enumerate(linha):
                if n % 2:
                    _sombrear(cels[i], CINZA_CLARO)
                par = cels[i].paragraphs[0]
                par.alignment = WD_ALIGN_PARAGRAPH.LEFT
                par.paragraph_format.line_spacing = 1.0
                self._escrever(par, str(valor), tamanho=TABELA_PT)
        if larguras_cm:
            for linha in tabela.rows:
                for i, largura in enumerate(larguras_cm):
                    linha.cells[i].width = Cm(largura)
        self.doc.add_paragraph()
        return tabela

    def fecho(self, local: str, data: str, formula: str = "Nestes termos,\nPede deferimento."):
        """Fórmula de encerramento, local/data e assinatura do advogado."""
        p = self.doc.add_paragraph()
        p.paragraph_format.first_line_indent = RECUO_PRIMEIRA_LINHA
        for i, linha in enumerate(formula.split("\n")):
            if i:
                p.add_run().add_break()
            self._escrever(p, linha)
        d = self.doc.add_paragraph()
        d.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._escrever(d, f"{local}, {data}.")
        a = self.doc.add_paragraph()
        a.alignment = WD_ALIGN_PARAGRAPH.CENTER
        a.paragraph_format.space_before = Pt(24)
        a.paragraph_format.keep_with_next = True
        self._escrever(a, ADVOGADO.upper(), negrito=True)
        o = self.doc.add_paragraph()
        o.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._escrever(o, OAB)
        return o

    # -- saída ----------------------------------------------------------
    def salvar(self, caminho: str) -> str:
        os.makedirs(os.path.dirname(os.path.abspath(caminho)), exist_ok=True)
        if os.path.exists(caminho):
            raise FileExistsError(f"Já existe: {caminho}. Escolha outro nome (nada é sobrescrito).")
        self.doc.save(caminho)
        return caminho


def docx_para_pdf(caminho_docx: str, pasta_saida: str | None = None) -> str:
    """Converte .docx em .pdf com o LibreOffice (precisa estar instalado). Mantém o layout do Word."""
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        raise RuntimeError("LibreOffice não encontrado. Instale-o (winget install TheDocumentFoundation.LibreOffice) "
                           "ou exporte o PDF pelo Word: Arquivo > Salvar como > PDF.")
    pasta = pasta_saida or os.path.dirname(os.path.abspath(caminho_docx))
    subprocess.run([exe, "--headless", "--convert-to", "pdf", "--outdir", pasta, caminho_docx],
                   check=True, capture_output=True, timeout=180)
    pdf = os.path.join(pasta, os.path.splitext(os.path.basename(caminho_docx))[0] + ".pdf")
    if not os.path.exists(pdf):
        raise RuntimeError("A conversão terminou sem gerar o PDF.")
    return pdf
